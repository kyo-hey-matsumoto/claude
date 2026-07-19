#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
DR.WU website-renewal-2026: 03/04/06 を1つのxlsx(複数タブ)にまとめてクライアント提出用に整形する。
"""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

BASE = "/Users/k.matsumoto/Desktop/US/claude_us/clients/drwu/website-renewal-2026"
OUT = "/private/tmp/claude-501/-Users-k-matsumoto-Desktop-US-claude-us/2ffe900c-d6bf-44fb-adf8-6bd6b8376014/scratchpad/DR.WU_website-renewal-2026_03-04-06.xlsx"

# ---------- palette (neutral / professional) ----------
HEADER_FILL = PatternFill("solid", fgColor="333333")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10, name="Meiryo UI")
TITLE_FILL = PatternFill("solid", fgColor="F2F2F2")
TITLE_FONT = Font(color="1A1A1A", bold=True, size=13, name="Meiryo UI")
SUBTITLE_FONT = Font(color="666666", size=9, italic=True, name="Meiryo UI")
BODY_FONT = Font(size=9.5, name="Meiryo UI")
BODY_FONT_B = Font(size=9.5, name="Meiryo UI", bold=True)
BAND_FILL = PatternFill("solid", fgColor="F7F7F7")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
WRAP_C = Alignment(wrap_text=True, vertical="center", horizontal="center")

# status-tag tint (very light, neutral-leaning)
TAG_FILLS = {
    "既存": PatternFill("solid", fgColor="EAF0F6"),       # pale blue-gray
    "新設": PatternFill("solid", fgColor="FBF3DE"),       # pale amber
    "レガシー": PatternFill("solid", fgColor="F7E6E4"),   # pale red
    "拡張": PatternFill("solid", fgColor="EDF3EA"),       # pale green-gray
    "刷新": PatternFill("solid", fgColor="FBF3DE"),
}
CLASS_FILLS = {
    "Shopify標準": PatternFill("solid", fgColor="EAF0F6"),
    "Shopifyアプリ": PatternFill("solid", fgColor="FBF3DE"),
    "独自開発": PatternFill("solid", fgColor="F1EAF6"),
}
PRIORITY_FONT = {
    "Must": Font(size=9.5, bold=True, color="1A1A1A", name="Meiryo UI"),
    "Should": Font(size=9.5, color="4D4D4D", name="Meiryo UI"),
    "Could": Font(size=9.5, color="808080", name="Meiryo UI"),
}


def tag_fill_for(value: str):
    if not value:
        return None
    for key, fill in TAG_FILLS.items():
        if key in value:
            return fill
    return None


def read_csv(path):
    with open(path, encoding="utf-8") as f:
        rows = list(csv.reader(f))
    return rows[0], rows[1:]


def setup_sheet(ws, title, subtitle, header, col_widths, freeze="A3"):
    ncols = len(header)
    last_col = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 26

    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[2].height = 16

    for i, h in enumerate(header, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = WRAP_C
    ws.row_dimensions[3].height = 28

    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    return 4  # next data row


def write_rows(ws, start_row, rows, col_widths, tag_col_idx=None, band=True,
               priority_col_idx=None, class_col_idx=None):
    r = start_row
    for row_i, row in enumerate(rows):
        fill = None
        if tag_col_idx is not None:
            fill = tag_fill_for(row[tag_col_idx])
        if class_col_idx is not None and fill is None:
            for key, f in CLASS_FILLS.items():
                if key in row[class_col_idx]:
                    fill = f
                    break
        if fill is None and band and row_i % 2 == 1:
            fill = BAND_FILL

        for ci, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=ci, value=val if val != "" else "")
            cell.border = BORDER
            cell.alignment = WRAP
            font = BODY_FONT
            if priority_col_idx is not None and ci - 1 == priority_col_idx:
                font = PRIORITY_FONT.get(val, BODY_FONT)
                cell.alignment = WRAP_C
            cell.font = font
            if fill:
                cell.fill = fill
        # rough row height based on longest cell text
        maxlen = max((len(str(v)) for v in row), default=0)
        ws.row_dimensions[r].height = max(16, min(90, 14 + maxlen // 10 * 4))
        r += 1
    return r


wb = Workbook()
wb.remove(wb.active)

# ============================================================
# Tab 1: 03_ディレクトリマップ
# ============================================================
header, rows = read_csv(f"{BASE}/03-directory-map.csv")
ws = wb.create_sheet("03_ディレクトリマップ")
col_widths = [16, 16, 18, 14, 30, 16, 10, 16, 34, 30, 30]
start = setup_sheet(
    ws,
    "03. ディレクトリマップ（drwu.jp 改修後サイトマップ）",
    "現行ページ階層の棚卸し＋改修意向を反映した改修後の構成（定期便/マイページ/フォーム一体型LP追加を含む）。第1階層セルの色は既存/新設/レガシーの区分を示す。",
    header, col_widths,
)
write_rows(ws, start, rows, col_widths, tag_col_idx=0)

# ============================================================
# Tab 2: 03_画面一覧
# ============================================================
header, rows = read_csv(f"{BASE}/03-screen-list.csv")
ws = wb.create_sheet("03_画面一覧")
col_widths = [8, 26, 14, 22, 26, 20, 14, 30]
start = setup_sheet(
    ws,
    "03. 画面一覧（URLではなく画面／テンプレート単位の整理）",
    "SCR-01〜SCR-26。区分列の色が既存/新設/レガシーに対応。「ワイヤーフレーム対象」列は05成果物との対応関係を示す（SCR-23〜26は追加候補・未着手）。",
    header, col_widths,
)
write_rows(ws, start, rows, col_widths, tag_col_idx=2)

# ============================================================
# Tab 3: 04_機能一覧
# ============================================================
header, rows = read_csv(f"{BASE}/04-feature-list.csv")
ws = wb.create_sheet("04_機能一覧")
col_widths = [8, 26, 18, 16, 16, 34, 34, 10, 26]
start = setup_sheet(
    ws,
    "04. 機能一覧（Shopify標準／アプリ／独自開発）",
    "F-01〜F-18。分類列の色がShopify標準/Shopifyアプリ/独自開発に対応。優先度はMust(濃)/Should/Could(淡)の階調で表現。",
    header, col_widths,
)
write_rows(ws, start, rows, col_widths, class_col_idx=2, priority_col_idx=7)

# ============================================================
# Tab 4: 06_外部連携・非機能要件
# ============================================================
ws = wb.create_sheet("06_外部連携・非機能要件")
ncols = 7
col_widths = [26, 24, 20, 16, 16, 26, 30]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.sheet_view.showGridLines = False

last_col = get_column_letter(ncols)
ws.merge_cells(f"A1:{last_col}1")
ws["A1"] = "06. 外部連携一覧・データ設計・非機能要件"
ws["A1"].font = TITLE_FONT
ws["A1"].fill = TITLE_FILL
ws["A1"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
ws.row_dimensions[1].height = 26
ws.merge_cells(f"A2:{last_col}2")
ws["A2"] = "LINE連携アプリ・バンドル/ランキングアプリ等の外部連携、想定データ設計、非機能要件の観点を整理"
ws["A2"].font = SUBTITLE_FONT
ws.row_dimensions[2].height = 16


def section_title(ws, row, text):
    ws.merge_cells(f"A{row}:{get_column_letter(ncols)}{row}")
    c = ws[f"A{row}"]
    c.value = text
    c.font = Font(bold=True, size=11, color="1A1A1A", name="Meiryo UI")
    c.fill = PatternFill("solid", fgColor="E4E4E4")
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 22
    return row + 1


def header_row(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = WRAP_C
    ws.row_dimensions[row].height = 24
    return row + 1


r = 4
r = section_title(ws, r, "① 外部連携一覧")
r = header_row(ws, r, ["連携先", "目的", "対応要件/機能", "データ方向", "認証方式", "想定コスト", "備考"])
integration_rows = [
    ["LINE公式アカウント + CRM PLUS on LINE", "LINE友だち連携、購入後自動配信、会員登録", "FR-08, FR-09 / F-09, F-10, F-11", "双方向（顧客属性・購入情報⇄LINE配信）", "LINE Login連携", "月額$100（〜20,000ID）／月額$200〜（ID連携無制限、ID数による追従課金あり）", "案A。ID数に応じた課金体系"],
    ["LINE公式アカウント + StoreCRM", "同上（代替案）", "同上", "同上", "同上", "月額$100〜（6,000ID以上は追従課金）", "案B。小規模スタートならコスト優位の可能性あり"],
    ["バンドル/セット販売アプリ（Bundles.app, PickyStory等）", "おすすめセットの商品化・在庫連動", "FR-04 / F-04", "Shopify商品・在庫データとの連携", "Shopify App連携（OAuth）", "アプリ次第（要選定）", "在庫共有方式かセット専用在庫方式かで設計が変わる"],
    ["ランキング集計アプリ（任意）", "売れ筋ランキングの自動集計", "FR-05 / F-05", "注文データの読み取り", "Shopify App連携（OAuth）", "アプリ次第（要選定）", "導入しない場合は運用者が月次で手動更新する代替案あり"],
    ["GA4（Google Analytics 4）", "TOPページ各セクションのクリック・CVR計測", "F-14", "サイト行動データの送信", "計測タグ（gtag.js）", "無料", "既存導入状況の確認が必要"],
    ["Shopify標準決済・配送", "購入・決済処理", "全体", "-", "Shopify標準", "Shopify利用料に含む", "変更なし"],
    ["ハックリベリー「定期購買」", "定期便（サブスクリプション）購入機能", "FR-12 / F-15, F-16", "双方向（商品・顧客・お届けスケジュール）", "Shopify App連携（OAuth）", "アプリ次第（要見積）", "案A。国内実績の多い定期購買アプリ"],
    ["Mikawaya Subscription", "定期便（サブスクリプション）購入機能（代替案）", "FR-12 / F-15, F-16", "同上", "Shopify App連携（OAuth）", "アプリ次第（要見積）", "案B。ハックリベリーと比較の上で選定"],
    ["シンプルフォーム一体型LP｜お手軽チャットボット購入（UnReact）", "フォーム一体型LPのチャットボット購入代替実装", "FR-14 / F-18", "入力データを標準決済ページへ自動反映", "Shopify App連携（テーマ拡張）", "月額$49.99〜（年払い17%割引）、7日間無料体験", "Shopifyはフォーム一体型LPを直接実装できないための代替案"],
]
r = write_rows(ws, r, integration_rows, col_widths, band=True) + 1

r = section_title(ws, r, "② データ設計（概要）")
data_design_lines = [
    "顧客(Customer) ── LINE友だちID（連携アプリ経由）／ポイント残高／注文履歴(Order)／会員ステータス（新規・継続・誕生日）",
    "商品(Product) ── シリーズタグ（マンデリック/ヒアルコンプレクス/…）／悩みタグ（毛穴/乾燥/くすみ/ハリ/敏感肌/UV、複数付与可）／セット構成（バンドルアプリ側で単品⇄セットを保持）",
    "キャンペーン(Campaign) ── 種別（初回限定/LINE特典/セット特典/季節企画）、掲載期間",
    "前提：現行商品には「シリーズ」分類はあるが「悩み」分類がないため、FR-03実装には商品タグの新規設計・全SKUへの付与が必要",
    "顧客(Customer) に定期便契約(Subscription)を追加 ── お届け周期・次回お届け日・スキップ/解約状態（定期便アプリ側で保持し、マイページから参照・操作）",
]
for line in data_design_lines:
    ws.merge_cells(f"A{r}:{last_col}{r}")
    c = ws[f"A{r}"]
    c.value = "・" + line
    c.font = BODY_FONT
    c.alignment = WRAP
    ws.row_dimensions[r].height = 18
    r += 1
r += 1

r = section_title(ws, r, "③ 非機能要件")
r = header_row(ws, r, ["区分", "要件", "備考", "", "", "", ""])
nfr_rows = [
    ["パフォーマンス", "TOPページのセクション数増加後も表示速度（Core Web Vitals）を現状同等以上に維持する", "画像の遅延読み込み・セクションの段階描画を検討", "", "", "", ""],
    ["計測", "TOP各セクション(①〜⑧)でのクリック率・コンバージョンをGA4イベントで取得できること", "F-14。改善のPDCAに必須", "", "", "", ""],
    ["SEO", "コレクション統廃合・URL変更時は301リダイレクトで既存の評価を継承する", "FR-11。レガシーコレクション整理時に必須", "", "", "", ""],
    ["セキュリティ・個人情報保護", "LINE連携・会員登録で扱う個人情報を適切に管理する", "プライバシーポリシーの記載更新を伴う", "", "", "", ""],
    ["可用性", "Shopifyプラットフォームおよび導入アプリのSLAに準拠する", "自社での特別な可用性対策は行わない", "", "", "", ""],
    ["保守性", "キャンペーンバナー・コラム・使い方ガイド等は運用担当者が更新できること", "Shopify標準ページ/メタフィールドでの運用が基本方針", "", "", "", ""],
    ["拡張性", "将来的な多言語対応・海外展開を著しく阻害しない設計とする", "対象外事項だが設計時に配慮", "", "", "", ""],
    ["アクセシビリティ", "主要導線はスマートフォンでの操作性を優先して設計する", "ターゲットユーザーの主要デバイスを想定", "", "", "", ""],
]
r = write_rows(ws, r, nfr_rows, col_widths, band=True) + 1

r = section_title(ws, r, "④ 未確定事項（要すり合わせ）")
open_items = [
    "LINE連携アプリの最終選定（CRM PLUS on LINE／StoreCRM）と、対象ID数に応じたコスト試算",
    "バンドル/セット販売アプリの選定と在庫連動方式",
    "ランキングの集計方式（アプリ自動集計／運用者による手動更新）",
    "GA4の導入状況・既存の計測設計の有無",
    "定期便アプリの最終選定（ハックリベリー「定期購買」／Mikawaya Subscription）と、顧客ポータル機能のマイページ埋め込み可否",
    "チャットボット購入アプリの最終選定（シンプルフォーム一体型LP／他のLPビルダー）と対象商品範囲（マンデリックのみか拡大するか）",
]
for line in open_items:
    ws.merge_cells(f"A{r}:{last_col}{r}")
    c = ws[f"A{r}"]
    c.value = "・" + line
    c.font = BODY_FONT
    c.alignment = WRAP
    ws.row_dimensions[r].height = 18
    r += 1

wb.save(OUT)
print("saved:", OUT)
